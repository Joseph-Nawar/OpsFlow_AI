from __future__ import annotations

import io
import socket
import struct
import warnings
from collections.abc import Callable, Iterable
from datetime import date, datetime, time
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl import load_workbook as openpyxl_load_workbook

from opsflow.documents import xlsx
from opsflow.documents.errors import (
    DocumentLimitError,
    DocumentParseError,
    DocumentValidationError,
)
from opsflow.documents.limits import DEFAULT_DOCUMENT_LIMITS, DocumentLimits
from opsflow.documents.models import CanonicalTable, DocumentWarning

XLSX_FIXTURE_ROOT = Path("fixtures/documents/xlsx")

OOXML_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OOXML_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
WORKSHEET_REL = f"{OOXML_REL}/worksheet"


def _limits(**overrides: int) -> DocumentLimits:
    from dataclasses import replace

    return replace(DEFAULT_DOCUMENT_LIMITS, **overrides)


def _sheet_xml(cells: str = "") -> bytes:
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="{OOXML_MAIN}">
  <dimension ref="A1:XFD1048576"/>
  <sheetData>{cells}</sheetData>
</worksheet>""".encode()


def _cell(coordinate: str, value: str, *, formula: str | None = None) -> str:
    formula_xml = f"<f>{formula}</f>" if formula is not None else ""
    row_number = coordinate.lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ")
    return f'<row r="{row_number}"><c r="{coordinate}">{formula_xml}<v>{value}</v></c></row>'


def _workbook_xml(sheet_count: int) -> bytes:
    sheets = "".join(
        f'<sheet name="Sheet{index}" sheetId="{index}" r:id="rId{index}"/>'
        for index in range(1, sheet_count + 1)
    )
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="{OOXML_MAIN}" xmlns:r="{OOXML_REL}">
  <sheets>{sheets}</sheets>
</workbook>""".encode()


def _relationships_xml(sheet_count: int) -> bytes:
    relationships = "".join(
        f'<Relationship Id="rId{index}" Type="{WORKSHEET_REL}" '
        f'Target="worksheets/sheet{index}.xml"/>'
        for index in range(1, sheet_count + 1)
    )
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="{PACKAGE_REL}">{relationships}</Relationships>""".encode()


def _content_types_xml(sheet_count: int, *, extra: str = "") -> bytes:
    overrides = "".join(
        f'<Override PartName="/xl/worksheets/sheet{index}.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for index in range(1, sheet_count + 1)
    )
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml"
            ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  {overrides}{extra}
</Types>""".encode()


def _valid_entries(
    *,
    sheet_contents: Iterable[bytes] = (_sheet_xml(),),
    workbook: bytes | None = None,
    relationships: bytes | None = None,
    content_types: bytes | None = None,
) -> list[tuple[str, bytes]]:
    sheets = list(sheet_contents)
    return [
        (
            "[Content_Types].xml",
            content_types or _content_types_xml(len(sheets)),
        ),
        ("xl/workbook.xml", workbook or _workbook_xml(len(sheets))),
        (
            "xl/_rels/workbook.xml.rels",
            relationships or _relationships_xml(len(sheets)),
        ),
        *[
            (f"xl/worksheets/sheet{index}.xml", content)
            for index, content in enumerate(sheets, start=1)
        ],
    ]


def _zip_bytes(entries: Iterable[tuple[str, bytes]]) -> bytes:
    output = io.BytesIO()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
            for name, content in entries:
                archive.writestr(name, content)
    return output.getvalue()


def _inflate_first_member_size(content: bytes, size: int) -> bytes:
    archive = bytearray(content)
    central_directory_signature = b"PK\x01\x02"
    offset = archive.find(central_directory_signature)
    assert offset >= 0
    struct.pack_into("<I", archive, offset + 24, size)
    return bytes(archive)


def test_xlsx_package_requires_content_types_and_workbook_parts(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(xlsx, "load_workbook", lambda *_args, **_kwargs: pytest.fail())

    for missing in ("[Content_Types].xml", "xl/workbook.xml"):
        entries = [entry for entry in _valid_entries() if entry[0] != missing]

        with pytest.raises(DocumentValidationError):
            xlsx._validate_xlsx_package(_zip_bytes(entries), DEFAULT_DOCUMENT_LIMITS)


def test_xlsx_package_rejects_non_zip_and_malformed_zip() -> None:
    valid = _zip_bytes(_valid_entries())

    with pytest.raises((DocumentParseError, DocumentValidationError)):
        xlsx._validate_xlsx_package(b"not a ZIP archive", DEFAULT_DOCUMENT_LIMITS)
    with pytest.raises((DocumentParseError, DocumentValidationError)):
        xlsx._validate_xlsx_package(valid[:-8], DEFAULT_DOCUMENT_LIMITS)


def test_xlsx_package_rejects_missing_or_malformed_relationship_metadata() -> None:
    missing = [entry for entry in _valid_entries() if entry[0] != "xl/_rels/workbook.xml.rels"]
    malformed = _valid_entries(relationships=b"<Relationships")

    with pytest.raises(DocumentValidationError):
        xlsx._validate_xlsx_package(_zip_bytes(missing), DEFAULT_DOCUMENT_LIMITS)
    with pytest.raises(DocumentValidationError):
        xlsx._validate_xlsx_package(_zip_bytes(malformed), DEFAULT_DOCUMENT_LIMITS)


def test_xlsx_package_rejects_duplicate_member_names() -> None:
    entries = _valid_entries()
    entries.append(("xl/workbook.xml", b"duplicate"))

    with pytest.raises(DocumentValidationError, match="duplicate"):
        xlsx._validate_xlsx_package(_zip_bytes(entries), DEFAULT_DOCUMENT_LIMITS)


def test_xlsx_package_rejects_macro_member_and_macro_content_type() -> None:
    macro_member = _valid_entries()
    macro_member.append(("xl/vbaProject.bin", b"macro bytes"))
    macro_type = _valid_entries(
        content_types=_content_types_xml(
            1,
            extra=(
                '<Override PartName="/xl/vbaProject.bin" '
                'ContentType="application/vnd.ms-office.vbaProject"/>'
            ),
        )
    )

    with pytest.raises(DocumentValidationError, match="macro"):
        xlsx._validate_xlsx_package(_zip_bytes(macro_member), DEFAULT_DOCUMENT_LIMITS)
    with pytest.raises(DocumentValidationError, match="macro"):
        xlsx._validate_xlsx_package(_zip_bytes(macro_type), DEFAULT_DOCUMENT_LIMITS)


def test_xlsx_package_enforces_expanded_size_before_workbook_load(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = _inflate_first_member_size(_zip_bytes(_valid_entries()), 101)
    monkeypatch.setattr(xlsx, "load_workbook", lambda *_args, **_kwargs: pytest.fail())

    with pytest.raises(DocumentLimitError, match="expanded"):
        xlsx._validate_xlsx_package(
            content,
            _limits(max_xlsx_expanded_bytes=100),
        )


def test_xlsx_package_rejects_sheet_count_before_workbook_materialization(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    entries = _valid_entries(sheet_contents=(_sheet_xml(), _sheet_xml()))
    monkeypatch.setattr(xlsx, "load_workbook", lambda *_args, **_kwargs: pytest.fail())

    with pytest.raises(DocumentLimitError, match="sheet"):
        xlsx._validate_xlsx_package(_zip_bytes(entries), _limits(max_xlsx_sheets=1))


def test_xlsx_package_sparse_row_guard_rejects_row_beyond_small_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = _zip_bytes(_valid_entries(sheet_contents=(_sheet_xml(_cell("A6", "1")),)))
    monkeypatch.setattr(xlsx, "load_workbook", lambda *_args, **_kwargs: pytest.fail())

    with pytest.raises(DocumentLimitError, match="row"):
        xlsx._validate_xlsx_package(
            content,
            _limits(max_xlsx_rows_per_sheet=5),
        )


def test_xlsx_package_sparse_column_guard_rejects_column_beyond_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = _zip_bytes(_valid_entries(sheet_contents=(_sheet_xml(_cell("C1", "1")),)))
    monkeypatch.setattr(xlsx, "load_workbook", lambda *_args, **_kwargs: pytest.fail())

    with pytest.raises(DocumentLimitError, match="column"):
        xlsx._validate_xlsx_package(
            content,
            _limits(max_table_columns=2),
        )


def test_xlsx_package_rejects_aggregate_populated_cell_limit() -> None:
    content = _zip_bytes(
        _valid_entries(sheet_contents=(_sheet_xml(_cell("A1", "1") + _cell("B1", "2")),))
    )

    with pytest.raises(DocumentLimitError, match="populated"):
        xlsx._validate_xlsx_package(
            content,
            _limits(max_xlsx_populated_cells=1),
        )


def test_xlsx_package_style_only_cells_do_not_expand_meaningful_bounds() -> None:
    style_only = '<row r="100"><c r="XFD100" s="1"/></row>'
    content = _zip_bytes(_valid_entries(sheet_contents=(_sheet_xml(style_only),)))

    package = xlsx._validate_xlsx_package(content, DEFAULT_DOCUMENT_LIMITS)

    assert package.populated_cell_count == 0
    assert package.sheets[0].max_meaningful_row == 0
    assert package.sheets[0].max_meaningful_column == 0


def test_xlsx_package_rejects_worksheet_relationship_outside_allowed_parts() -> None:
    relationships = _relationships_xml(1).replace(
        b'Target="worksheets/sheet1.xml"', b'Target="../external.xml"'
    )
    content = _zip_bytes(_valid_entries(relationships=relationships))

    with pytest.raises(DocumentValidationError, match="worksheet"):
        xlsx._validate_xlsx_package(content, DEFAULT_DOCUMENT_LIMITS)


def test_xlsx_package_rejects_malformed_cell_coordinate_safely() -> None:
    content = _zip_bytes(_valid_entries(sheet_contents=(_sheet_xml(_cell("A0", "1")),)))

    with pytest.raises(DocumentValidationError, match="coordinate"):
        xlsx._validate_xlsx_package(content, DEFAULT_DOCUMENT_LIMITS)


def _workbook_bytes(configure: Callable[[Workbook], object]) -> bytes:
    workbook = Workbook()
    try:
        configure(workbook)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def _set_cell(workbook: Workbook, coordinate: str, value: object) -> None:
    workbook.active[coordinate] = value


def _add_external_link(content: bytes) -> bytes:
    external_link_relationship = (
        b'<Relationship Id="rIdExternal" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" '
        b'Target="externalLinks/externalLink1.xml"/>'
    )
    entries: list[tuple[str, bytes]] = []
    with ZipFile(io.BytesIO(content)) as archive:
        for info in archive.infolist():
            member = archive.read(info.filename)
            if info.filename == "xl/_rels/workbook.xml.rels":
                member = member.replace(
                    b"</Relationships>",
                    external_link_relationship + b"</Relationships>",
                )
            entries.append((info.filename, member))
    entries.extend(
        [
            (
                "xl/externalLinks/externalLink1.xml",
                f'<externalLink xmlns="{OOXML_MAIN}"/>'.encode(),
            ),
            (
                "xl/externalLinks/_rels/externalLink1.xml.rels",
                (
                    f'<Relationships xmlns="{PACKAGE_REL}">'
                    '<Relationship Id="rIdPath" '
                    'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    'relationships/externalLinkPath" '
                    'Target="http://example.test/external.xlsx" TargetMode="External"/>'
                    "</Relationships>"
                ).encode(),
            ),
        ]
    )
    return _zip_bytes(entries)


def test_parse_xlsx_document_preserves_sheet_order_and_names() -> None:
    content = (XLSX_FIXTURE_ROOT / "clean-multisheet.xlsx").read_bytes()

    parsed = xlsx.parse_xlsx_document(content, DEFAULT_DOCUMENT_LIMITS)

    assert tuple(table.name for table in parsed.tables) == ("Orders", "Notes")
    assert parsed.tables[0].rows == (("SKU", "Quantity"), ("ABC-1", "10"))
    assert parsed.tables[1].rows == (("Comment",), ("Synthetic fixture",))
    assert parsed.text == (
        '# Sheet: "Orders"\n["SKU","Quantity"]\n["ABC-1","10"]\n\n'
        '# Sheet: "Notes"\n["Comment"]\n["Synthetic fixture"]'
    )
    assert parsed.pages == ()


def test_parse_xlsx_document_preserves_structural_empty_cells_and_rows() -> None:
    content = _workbook_bytes(
        lambda workbook: (
            setattr(workbook.active, "title", "Sparse"),
            _set_cell(workbook, "A1", "A1"),
            _set_cell(workbook, "C2", "C2"),
        )
    )

    parsed = xlsx.parse_xlsx_document(content, DEFAULT_DOCUMENT_LIMITS)

    assert parsed.tables == (CanonicalTable(name="Sparse", rows=(("A1", "", ""), ("", "", "C2"))),)


def test_parse_xlsx_document_preserves_formula_expression() -> None:
    content = (XLSX_FIXTURE_ROOT / "formulas.xlsx").read_bytes()

    parsed = xlsx.parse_xlsx_document(content, DEFAULT_DOCUMENT_LIMITS)

    assert parsed.tables[0].rows[2][0] == "=SUM(A1:A2)"
    assert "3" not in parsed.tables[0].rows[2]


def test_parse_xlsx_document_converts_scalars_deterministically() -> None:
    content = _workbook_bytes(
        lambda workbook: (
            setattr(workbook.active, "title", "Scalars"),
            workbook.active.append(
                [
                    "Cafe\u0301",
                    None,
                    False,
                    True,
                    7,
                    1.5,
                    date(2026, 9, 15),
                    datetime(2026, 9, 15, 10, 11, 12),
                    time(13, 14, 15),
                ]
            ),
        )
    )

    parsed = xlsx.parse_xlsx_document(content, DEFAULT_DOCUMENT_LIMITS)

    assert parsed.tables[0].rows == (
        (
            "Café",
            "",
            "false",
            "true",
            "7",
            "1.5",
            "2026-09-15T00:00:00",
            "2026-09-15T10:11:12",
            "13:14:15",
        ),
    )


def test_parse_xlsx_document_warns_only_for_truly_empty_sheet() -> None:
    content = _workbook_bytes(
        lambda workbook: (
            setattr(workbook.active, "title", "Empty"),
            workbook.active["A100"].__setattr__("style", "Normal"),
        )
    )

    parsed = xlsx.parse_xlsx_document(content, DEFAULT_DOCUMENT_LIMITS)

    assert parsed.tables == (CanonicalTable(name="Empty", rows=()),)
    assert parsed.warnings == (
        DocumentWarning(
            code="EMPTY_SHEET",
            message="Worksheet contains no non-blank cells.",
            location="sheet:Empty",
        ),
    )


def test_parse_xlsx_document_enforces_sheet_row_cell_and_column_limits() -> None:
    two_sheets = _workbook_bytes(lambda workbook: workbook.create_sheet("Second"))
    with pytest.raises(DocumentLimitError, match="sheet"):
        xlsx.parse_xlsx_document(two_sheets, _limits(max_xlsx_sheets=1))

    row_six = _workbook_bytes(lambda workbook: _set_cell(workbook, "A6", "value"))
    with pytest.raises(DocumentLimitError, match="row"):
        xlsx.parse_xlsx_document(row_six, _limits(max_xlsx_rows_per_sheet=5))

    column_c = _workbook_bytes(lambda workbook: _set_cell(workbook, "C1", "value"))
    with pytest.raises(DocumentLimitError, match="column"):
        xlsx.parse_xlsx_document(column_c, _limits(max_table_columns=2))

    two_cells = _workbook_bytes(
        lambda workbook: (
            _set_cell(workbook, "A1", "one"),
            _set_cell(workbook, "B1", "two"),
        )
    )
    with pytest.raises(DocumentLimitError, match="populated"):
        xlsx.parse_xlsx_document(two_cells, _limits(max_xlsx_populated_cells=1))


def test_parse_xlsx_document_preserves_per_sheet_row_span_without_sparse_workload() -> None:
    row_five = _workbook_bytes(lambda workbook: _set_cell(workbook, "A5", "value"))

    parsed = xlsx.parse_xlsx_document(
        row_five,
        _limits(max_xlsx_rows_per_sheet=5),
    )

    assert parsed.tables[0].rows == (("",), ("",), ("",), ("",), ("value",))


def test_parse_xlsx_document_rejects_unexpected_scalar_type(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakeCell:
        value = object()

    class FakeWorksheet:
        title = "Fake"

        def iter_rows(self, **_kwargs: int) -> Iterable[tuple[FakeCell]]:
            yield (FakeCell(),)

    class FakeWorkbook:
        worksheets = (FakeWorksheet(),)
        closed = False

        def close(self) -> None:
            self.closed = True

    package = xlsx._XlsxPackageInfo(
        sheets=(
            xlsx._XlsxSheetBounds(
                title="Fake",
                part_name="xl/worksheets/sheet1.xml",
                max_meaningful_row=1,
                max_meaningful_column=1,
            ),
        ),
        populated_cell_count=1,
    )
    workbook = FakeWorkbook()
    monkeypatch.setattr(xlsx, "_validate_xlsx_package", lambda *_args: package)
    monkeypatch.setattr(xlsx, "load_workbook", lambda *_args, **_kwargs: workbook)

    with pytest.raises(DocumentParseError, match="scalar"):
        xlsx.parse_xlsx_document(b"ignored", DEFAULT_DOCUMENT_LIMITS)

    assert workbook.closed is True


def test_parse_xlsx_document_does_not_follow_external_links(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = _add_external_link(
        _workbook_bytes(lambda workbook: _set_cell(workbook, "A1", "local"))
    )
    calls: list[dict[str, object]] = []

    def fail_network(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("network access is forbidden")

    def recording_loader(*args: object, **kwargs: object) -> object:
        calls.append(kwargs)
        return openpyxl_load_workbook(*args, **kwargs)

    monkeypatch.setattr(xlsx, "load_workbook", recording_loader)
    monkeypatch.setattr(socket, "create_connection", fail_network)

    parsed = xlsx.parse_xlsx_document(content, DEFAULT_DOCUMENT_LIMITS)

    assert parsed.tables[0].rows == (("local",),)
    assert calls[0]["read_only"] is True
    assert calls[0]["data_only"] is False
    assert calls[0]["keep_links"] is False


def test_parse_xlsx_document_rejects_corrupt_workbook_and_non_xlsx_zip() -> None:
    corrupt_workbook = _zip_bytes(_valid_entries(workbook=b"not XML"))
    arbitrary_zip = _zip_bytes((("payload.bin", b"not an XLSX"),))

    with pytest.raises((DocumentParseError, DocumentValidationError)):
        xlsx.parse_xlsx_document(corrupt_workbook, DEFAULT_DOCUMENT_LIMITS)
    with pytest.raises((DocumentParseError, DocumentValidationError)):
        xlsx.parse_xlsx_document(arbitrary_zip, DEFAULT_DOCUMENT_LIMITS)


def test_parse_xlsx_document_repeated_parses_compare_equal() -> None:
    content = (XLSX_FIXTURE_ROOT / "clean-multisheet.xlsx").read_bytes()

    first = xlsx.parse_xlsx_document(content, DEFAULT_DOCUMENT_LIMITS)
    second = xlsx.parse_xlsx_document(content, DEFAULT_DOCUMENT_LIMITS)

    assert first == second
