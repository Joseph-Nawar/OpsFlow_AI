"""Behavioral tests for the Phase 3 unified document processor."""

import hashlib
import io
import subprocess
import sys
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook

import opsflow.documents.processor as processor
from opsflow.documents.common import sha256_bytes
from opsflow.documents.errors import (
    DocumentLimitError,
    DocumentParseError,
    DocumentValidationError,
    UnsupportedDocumentTypeError,
)
from opsflow.documents.limits import DEFAULT_DOCUMENT_LIMITS, DocumentLimits
from opsflow.documents.models import DocumentInput, _ParsedDocumentContent
from opsflow.domain.records import SourceDocumentType

FIXTURE_ROOT = Path("fixtures/documents")
MIME_BY_TYPE = {
    SourceDocumentType.EMAIL_BODY: "text/plain",
    SourceDocumentType.CSV: "text/csv",
    SourceDocumentType.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    SourceDocumentType.PDF: "application/pdf",
}


def _limits(**overrides: int) -> DocumentLimits:
    return replace(DEFAULT_DOCUMENT_LIMITS, **overrides)


def _document(
    document_type: SourceDocumentType,
    content: bytes,
    *,
    name: str = "document.bin",
    mime_type: str | None = None,
    source_reference: str | None = None,
    metadata: tuple[tuple[str, str], ...] = (),
) -> DocumentInput:
    return DocumentInput(
        document_type=document_type,
        name=name,
        mime_type=mime_type or MIME_BY_TYPE[document_type],
        content=content,
        source_reference=source_reference,
        metadata=metadata,
    )


def _fixture(path: str) -> bytes:
    return (FIXTURE_ROOT / path).read_bytes()


def _xlsx_with_cell(coordinate: str, value: object) -> bytes:
    workbook = Workbook()
    try:
        workbook.active[coordinate] = value
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def test_process_document_builds_common_envelope_for_plain_text() -> None:
    content = b"SKU,Quantity\r\nABC-1,2"
    document = _document(
        SourceDocumentType.EMAIL_BODY,
        content,
        name="order.pdf",
        mime_type=" Text/Plain ; charset=utf-8 ",
        source_reference="message-42",
        metadata=(("tag", "one"), ("tag", "two")),
    )

    result = processor.process_document(document)

    assert result.document_type is SourceDocumentType.EMAIL_BODY
    assert result.name == "order.pdf"
    assert result.mime_type == "text/plain"
    assert result.sha256 == hashlib.sha256(content).hexdigest()
    assert result.size_bytes == len(content)
    assert result.source_reference == "message-42"
    assert result.metadata == (("tag", "one"), ("tag", "two"))
    assert result.text == "SKU,Quantity\nABC-1,2"
    assert result.pages == ()
    assert result.tables == ()
    assert result.warnings == ()


def test_process_document_dispatches_every_supported_format() -> None:
    inputs = (
        _document(
            SourceDocumentType.EMAIL_BODY,
            _fixture("text/clean-po.txt"),
            name="body.txt",
        ),
        _document(
            SourceDocumentType.CSV,
            _fixture("csv/quoted-multiline.csv"),
            name="orders.csv",
        ),
        _document(
            SourceDocumentType.XLSX,
            _fixture("xlsx/clean-multisheet.xlsx"),
            name="orders.xlsx",
        ),
        _document(
            SourceDocumentType.PDF,
            _fixture("pdf/single-page-text.pdf"),
            name="orders.pdf",
        ),
    )

    results = tuple(processor.process_document(document) for document in inputs)

    assert results[0].text
    assert len(results[1].tables) == 1
    assert len(results[2].tables) == 2
    assert len(results[3].pages) == 1


def test_process_document_rejects_form_before_parser_dispatch(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_parser(*_args: object, **_kwargs: object) -> _ParsedDocumentContent:
        pytest.fail("FORM reached a format parser")

    for parser_name in (
        "parse_text_document",
        "parse_csv_document",
        "parse_xlsx_document",
        "parse_pdf_document",
    ):
        monkeypatch.setattr(processor, parser_name, fail_parser)

    document = _document(
        SourceDocumentType.FORM,
        b"form source",
        name="form.any",
        mime_type="application/octet-stream",
    )

    with pytest.raises(UnsupportedDocumentTypeError):
        processor.process_document(document)


@pytest.mark.parametrize(
    ("document_type", "mime_type"),
    [
        (
            SourceDocumentType.PDF,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        (SourceDocumentType.XLSX, "application/pdf"),
        (SourceDocumentType.CSV, "text/plain"),
    ],
)
def test_process_document_enforces_exact_mime_map_without_reclassification(
    document_type: SourceDocumentType,
    mime_type: str,
) -> None:
    with pytest.raises(DocumentValidationError, match="MIME"):
        processor.process_document(
            _document(
                document_type,
                b"content",
                name="document.with-wrong-extension",
                mime_type=mime_type,
            )
        )


def test_process_document_does_not_use_filename_extension_for_type_detection() -> None:
    result = processor.process_document(
        _document(
            SourceDocumentType.EMAIL_BODY,
            b"ordinary body",
            name="ordinary.pdf",
            mime_type="text/plain; charset=utf-8",
        )
    )

    assert result.document_type is SourceDocumentType.EMAIL_BODY
    assert result.text == "ordinary body"


def test_process_document_rejects_oversized_input_before_parser(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_parser(*_args: object, **_kwargs: object) -> _ParsedDocumentContent:
        pytest.fail("oversized input reached a parser")

    for parser_name in (
        "parse_text_document",
        "parse_csv_document",
        "parse_xlsx_document",
        "parse_pdf_document",
    ):
        monkeypatch.setattr(processor, parser_name, fail_parser)

    with pytest.raises(DocumentLimitError, match="input"):
        processor.process_document(
            _document(SourceDocumentType.EMAIL_BODY, b"123456"),
            _limits(max_input_bytes=5),
        )


def test_process_document_rejects_canonical_text_overflow_without_truncation(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        processor,
        "parse_text_document",
        lambda _content: _ParsedDocumentContent(text="12345"),
    )

    with pytest.raises(DocumentLimitError, match="canonical text"):
        processor.process_document(
            _document(SourceDocumentType.EMAIL_BODY, b"source"),
            _limits(max_text_characters=4),
        )


def test_process_document_accepts_exact_input_and_text_limit_boundaries() -> None:
    content = b"body"

    result = processor.process_document(
        _document(SourceDocumentType.EMAIL_BODY, content),
        _limits(max_input_bytes=len(content), max_text_characters=len("body")),
    )

    assert result.text == "body"


def test_process_document_repeated_same_input_compares_equal() -> None:
    document = _document(
        SourceDocumentType.CSV,
        _fixture("csv/empty-cells.csv"),
        name="orders.csv",
        source_reference="source-1",
        metadata=(("tag", "one"), ("tag", "two")),
    )

    first = processor.process_document(document)
    second = processor.process_document(document)

    assert first == second


def test_process_document_hashes_exact_bytes_not_normalized_text() -> None:
    content = b"\xef\xbb\xbfCafe\r\n"
    document = _document(SourceDocumentType.EMAIL_BODY, content)

    result = processor.process_document(document)

    assert result.sha256 == sha256_bytes(content)
    assert result.sha256 != sha256_bytes("Café\n".encode())


def test_process_document_preserves_metadata_order_and_duplicate_keys() -> None:
    metadata = (("source", "mailbox"), ("tag", "one"), ("tag", "two"))
    document = _document(
        SourceDocumentType.EMAIL_BODY,
        b"body",
        source_reference="reference-1",
        metadata=metadata,
    )

    result = processor.process_document(document)

    assert result.metadata == metadata
    assert result.source_reference == "reference-1"


def test_supported_formats_have_canonical_output_without_infrastructure() -> None:
    script = r"""
import sys
from pathlib import Path

for module_name in (
    "opsflow.api",
    "opsflow.application",
    "opsflow.database",
    "opsflow.persistence",
    "fastapi",
    "httpx",
    "openai",
    "sqlalchemy",
):
    sys.modules[module_name] = None

from opsflow.documents import DocumentInput, process_document
from opsflow.domain.records import SourceDocumentType

root = Path("fixtures/documents")
inputs = (
    DocumentInput(
        SourceDocumentType.EMAIL_BODY,
        "body.txt",
        "text/plain",
        (root / "text/clean-po.txt").read_bytes(),
    ),
    DocumentInput(
        SourceDocumentType.CSV,
        "orders.csv",
        "text/csv",
        (root / "csv/empty-cells.csv").read_bytes(),
    ),
    DocumentInput(
        SourceDocumentType.XLSX,
        "orders.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        (root / "xlsx/clean-multisheet.xlsx").read_bytes(),
    ),
    DocumentInput(
        SourceDocumentType.PDF,
        "orders.pdf",
        "application/pdf",
        (root / "pdf/single-page-text.pdf").read_bytes(),
    ),
)
results = tuple(process_document(document) for document in inputs)
assert all(result.sha256 and result.size_bytes > 0 for result in results)
"""
    completed = subprocess.run(
        [sys.executable, "-c", script],
        capture_output=True,
        cwd=Path.cwd(),
        text=True,
    )

    assert completed.returncode == 0, completed.stderr


def test_document_package_import_has_no_forbidden_infrastructure_boundary() -> None:
    script = r"""
import sys
import opsflow.documents

for module_name in (
    "opsflow.api",
    "opsflow.application",
    "opsflow.database",
    "opsflow.persistence",
    "fastapi",
    "httpx",
    "openai",
    "sqlalchemy",
):
    assert module_name not in sys.modules, module_name
"""
    completed = subprocess.run(
        [sys.executable, "-c", script],
        capture_output=True,
        cwd=Path.cwd(),
        text=True,
    )

    assert completed.returncode == 0, completed.stderr


@pytest.mark.parametrize(
    ("document_type", "mime_type", "content", "expected_error"),
    [
        (SourceDocumentType.EMAIL_BODY, "text/plain", b"\xff", DocumentParseError),
        (SourceDocumentType.CSV, "text/csv", b'"unterminated', DocumentParseError),
        (
            SourceDocumentType.XLSX,
            MIME_BY_TYPE[SourceDocumentType.XLSX],
            b"not a ZIP",
            DocumentParseError,
        ),
        (SourceDocumentType.PDF, "application/pdf", b"%PDF-1.7\ntruncated", DocumentParseError),
    ],
)
def test_corrupt_documents_fail_with_safe_document_errors(
    document_type: SourceDocumentType,
    mime_type: str,
    content: bytes,
    expected_error: type[DocumentParseError],
) -> None:
    with pytest.raises(expected_error) as error:
        processor.process_document(
            _document(document_type, content, mime_type=mime_type),
        )

    assert content.decode("utf-8", errors="replace") not in str(error.value)


def test_processor_propagates_mime_mismatch_without_reclassification() -> None:
    with pytest.raises(DocumentValidationError, match="MIME"):
        processor.process_document(
            _document(
                SourceDocumentType.PDF,
                _fixture("pdf/single-page-text.pdf"),
                mime_type="text/plain",
                name="document.txt",
            )
        )


def test_processor_enforces_each_configured_limit_without_truncation() -> None:
    with pytest.raises(DocumentLimitError, match="input"):
        processor.process_document(
            _document(SourceDocumentType.EMAIL_BODY, b"1234"),
            _limits(max_input_bytes=3),
        )

    with pytest.raises(DocumentLimitError, match="canonical text"):
        processor.process_document(
            _document(SourceDocumentType.EMAIL_BODY, b"1234"),
            _limits(max_text_characters=3),
        )

    with pytest.raises(DocumentLimitError, match="PDF page"):
        processor.process_document(
            _document(SourceDocumentType.PDF, _fixture("pdf/multi-page-text.pdf")),
            _limits(max_pdf_pages=1),
        )

    with pytest.raises(DocumentLimitError, match="sheet"):
        processor.process_document(
            _document(SourceDocumentType.XLSX, _fixture("xlsx/clean-multisheet.xlsx")),
            _limits(max_xlsx_sheets=1),
        )

    with pytest.raises(DocumentLimitError, match="row"):
        processor.process_document(
            _document(SourceDocumentType.XLSX, _xlsx_with_cell("A6", "value")),
            _limits(max_xlsx_rows_per_sheet=5),
        )

    with pytest.raises(DocumentLimitError, match="populated"):
        processor.process_document(
            _document(SourceDocumentType.XLSX, _fixture("xlsx/clean-multisheet.xlsx")),
            _limits(max_xlsx_populated_cells=1),
        )

    with pytest.raises(DocumentLimitError, match="column"):
        processor.process_document(
            _document(SourceDocumentType.CSV, b"A,B,C\n1,2,3"),
            _limits(max_table_columns=2),
        )

    with pytest.raises(DocumentLimitError, match="expanded"):
        processor.process_document(
            _document(SourceDocumentType.XLSX, _fixture("xlsx/clean-multisheet.xlsx")),
            _limits(max_xlsx_expanded_bytes=1),
        )


def test_duplicate_hashes_are_exposed_but_not_rejected() -> None:
    content = _fixture("text/clean-po.txt")
    first = processor.process_document(
        _document(
            SourceDocumentType.EMAIL_BODY,
            content,
            name="first.txt",
            source_reference="reference-one",
            metadata=(("tag", "one"),),
        )
    )
    second = processor.process_document(
        _document(
            SourceDocumentType.EMAIL_BODY,
            content,
            name="second.txt",
            source_reference="reference-two",
            metadata=(("tag", "two"),),
        )
    )

    assert first.sha256 == second.sha256
    assert first.name == "first.txt"
    assert second.name == "second.txt"
    assert first.source_reference != second.source_reference
    assert first.metadata != second.metadata


def test_canonical_output_contains_no_runtime_identity() -> None:
    result = processor.process_document(
        _document(
            SourceDocumentType.EMAIL_BODY,
            b"stable body",
            name="stable.txt",
            source_reference="stable-reference",
        )
    )

    assert set(result.__dataclass_fields__) == {
        "document_type",
        "name",
        "mime_type",
        "sha256",
        "size_bytes",
        "source_reference",
        "metadata",
        "text",
        "pages",
        "tables",
        "warnings",
    }
    assert str(Path.cwd()) not in repr(result)
    assert "timestamp" not in repr(result).lower()
    assert "uuid" not in repr(result).lower()


def test_full_canonical_output_is_repeatable_for_all_supported_formats() -> None:
    documents = (
        _document(SourceDocumentType.EMAIL_BODY, _fixture("text/clean-po.txt")),
        _document(SourceDocumentType.CSV, _fixture("csv/empty-cells.csv")),
        _document(SourceDocumentType.XLSX, _fixture("xlsx/clean-multisheet.xlsx")),
        _document(SourceDocumentType.PDF, _fixture("pdf/multi-page-text.pdf")),
    )

    for document in documents:
        first = processor.process_document(document)
        second = processor.process_document(document)
        assert first == second
